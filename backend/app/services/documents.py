"""This file handles document uploads and indexing in Aether AI. It accepts PDF, Word, text, and Markdown files, extracts their text, splits it into smaller chunks, stores the document in the database, and indexes the chunks for AI search (RAG)."""


# clean_text()	Removes unwanted characters and extra spaces while preserving paragraphs.
# extract_text()	Extracts readable text from PDF, DOCX, TXT, and Markdown files. uses clean_text()
# chunk_text()	Splits extracted text into smaller chunks of 800 characters, with 150 characters of overlap.
# upload_document()	Main upload function. Validates files, detects duplicates, saves the document, and starts indexing.
# index_document()	Indexes document chunks in the vector store and updates the database with indexing status.
import base64
import hashlib
import logging
import re
from io import BytesIO
from pathlib import Path
from uuid import uuid4

import pymupdf
from docx import Document as WordDocument
from langchain_text_splitters import RecursiveCharacterTextSplitter
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from app.core.config import config
from app.core.memory import check_and_collect_garbage, log_stage_memory
from app.models.database import Document, new_id
from app.services import storage

_ocr_routing_store = None


def configure_ocr_router(store):
    global _ocr_routing_store
    _ocr_routing_store = store


SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md", ".markdown"}


def clean_text(text: str) -> str:
    """Remove nulls and repeated spaces while preserving paragraph boundaries."""
    text = text.replace("\x00", "").replace("\r\n", "\n")
    return re.sub(r"\n{3,}", "\n\n", re.sub(r"[ \t]+", " ", text)).strip()


def ocr_image_bytes(image_bytes: bytes, mime_type: str = "image/jpeg") -> str:
    """Transcribe readable text from image bytes using Gemini Vision with model fallbacks."""
    logger = logging.getLogger("aether.documents")
    try:
        b64 = base64.b64encode(image_bytes).decode("ascii")
    except Exception as exc:
        logger.warning("OCR: failed to base64 encode image: %s", exc)
        return ""

    from langchain_core.messages import HumanMessage
    from app.services.ai import create_llm

    prompt_msg = HumanMessage(
        content=[
            {
                "type": "text",
                "text": (
                    "Transcribe all readable text from this scanned document or image accurately. "
                    "Preserve original paragraph structure. "
                    "Return ONLY the transcribed text without conversational commentary or markdown code blocks."
                ),
            },
            {
                "type": "image_url",
                "image_url": {"url": f"data:{mime_type};base64,{b64}"},
            },
        ]
    )

    from app.services.ai import IncompleteAnswer, call_with_free_fallback, message_text, safe_service_error
    from app.services.free_models import get_free_models_sync
    import asyncio
    from time import time

    deadline = time() + 75
    try:
        catalog = get_free_models_sync()
        models = [item["id"] for item in catalog if item.get("available", True) and "vision" in item.get("supported_parameters", [])]
        if config.gemini_ocr_model in models:
            models.remove(config.gemini_ocr_model)
            models.insert(0, config.gemini_ocr_model)
        plan = {"models": models, "catalog": catalog, "unavailable": [], "deadline_at": deadline}

        async def operation(model):
            response = await asyncio.to_thread(create_llm(model).invoke, [prompt_msg])
            if getattr(response, "response_metadata", {}).get("finish_reason") == "MAX_TOKENS":
                raise IncompleteAnswer("OCR output was truncated.")
            text = clean_text(message_text(response))
            if not text:
                raise IncompleteAnswer("OCR returned no readable text.")
            return text

        return asyncio.run(call_with_free_fallback(
            plan, operation, lambda failed, replacement, reason: logger.warning("OCR fallback model=%s next=%s reason=%s", failed, replacement, reason),
            store=_ocr_routing_store, task="ocr", input_tokens=len(b64) + 256,
        ))
    except Exception as error:
        logger.warning("OCR could not complete: %s", safe_service_error(error))
        return ""


def ocr_page_image(page) -> str:
    """Render a scanned PDF page to JPEG and transcribe using Gemini Vision."""
    logger = logging.getLogger("aether.documents")
    try:
        pix = page.get_pixmap(dpi=150)
        jpeg_bytes = pix.tobytes("jpeg")
        del pix
        return ocr_image_bytes(jpeg_bytes, mime_type="image/jpeg")
    except Exception as exc:
        logger.warning("OCR: failed to render page to image: %s", exc)
        return ""


def extract_text(source: bytes | Path | str, extension: str) -> list[dict]: 
    """Extract page-aware PDF text, Word paragraphs/tables, or UTF-8 plain text.

    Supports both file paths (for low memory usage) and bytes (for in-memory sources).
    """
    max_chars = config.max_extracted_chars
    total_chars = 0
    segments = []

    try:
        if extension == ".pdf":
            pdf_context = (
                pymupdf.open(str(source))
                if isinstance(source, (str, Path))
                else pymupdf.open(stream=source, filetype="pdf")
            )
            with pdf_context as pdf:
                if pdf.is_encrypted:
                    raise ValueError("Password-protected PDFs are not supported.")
                for i, page in enumerate(pdf):
                    cleaned = clean_text(page.get_text())
                    if not cleaned and config.gemini_ocr_enabled and config.gemini_api_key:
                        cleaned = ocr_page_image(page)
                    if cleaned:
                        total_chars += len(cleaned)
                        if total_chars > max_chars:
                            raise ValueError(
                                f"Extracted text exceeds the {max_chars // 1_000_000} million character limit. Split this document."
                            )
                        segments.append({"text": cleaned, "page": i + 1})

        elif extension == ".docx":
            doc = (
                WordDocument(str(source))
                if isinstance(source, (str, Path))
                else WordDocument(BytesIO(source))
            )
            paragraphs = [paragraph.text for paragraph in doc.paragraphs]
            paragraphs += [
                " | ".join(cell.text for cell in row.cells) for table in doc.tables for row in table.rows
            ]
            cleaned = clean_text("\n\n".join(paragraphs))

            # If no direct text was found, inspect embedded images (e.g. scanned documents pasted into Word)
            if not cleaned and config.gemini_ocr_enabled and config.gemini_api_key:
                image_texts = []
                for rel in doc.part.rels.values():
                    if "image" in getattr(rel, "reltype", ""):
                        try:
                            image_part = rel.target_part
                            img_bytes = image_part.blob
                            img_mime = getattr(image_part, "content_type", "") or "image/jpeg"
                            img_text = ocr_image_bytes(img_bytes, mime_type=img_mime)
                            if img_text:
                                image_texts.append(img_text)
                        except Exception as exc:
                            logging.getLogger("aether.documents").warning("OCR failed for docx image: %s", exc)
                if image_texts:
                    cleaned = clean_text("\n\n".join(image_texts))

            if cleaned:
                if len(cleaned) > max_chars:
                    raise ValueError(
                        f"Extracted text exceeds the {max_chars // 1_000_000} million character limit. Split this document."
                    )
                segments.append({"text": cleaned, "page": 0})

        elif extension in {".txt", ".md", ".markdown"}:
            raw_text = (
                Path(source).read_text(encoding="utf-8-sig")
                if isinstance(source, (str, Path))
                else source.decode("utf-8-sig")
            )
            cleaned = clean_text(raw_text)
            if cleaned:
                if len(cleaned) > max_chars:
                    raise ValueError(
                        f"Extracted text exceeds the {max_chars // 1_000_000} million character limit. Split this document."
                    )
                segments.append({"text": cleaned, "page": 0})
        else:
            raise ValueError("Supported files: PDF, DOCX, TXT, and Markdown.")

    except (ValueError, UnicodeDecodeError) as error:
        raise ValueError(
            str(error) if isinstance(error, ValueError) else "Save text files as UTF-8."
        ) from error
    except Exception as error:
        raise ValueError("This file is damaged or its contents do not match its extension.") from error

    if not segments:
        if extension == ".pdf":
            if not (config.gemini_ocr_enabled and config.gemini_api_key):
                raise ValueError(
                    "No readable text found in this PDF. If this is a scanned document, ensure GEMINI_API_KEY is configured in backend environment for automatic OCR."
                )
            raise ValueError(
                "No readable text found in this PDF. Please ensure the document is not blank and pages are clearly legible."
            )
        elif extension == ".docx":
            raise ValueError(
                "No readable text found in this Word document. Please ensure the file contains readable text or clear images."
            )
        else:
            raise ValueError(f"No readable text found in this {extension[1:].upper()} file.")

    return segments


def chunk_text(segments: list[dict]) -> list[dict]:
    splitter = RecursiveCharacterTextSplitter(chunk_size=800, chunk_overlap=150)
    return [
        {"text": chunk, "page": segment["page"]}
        for segment in segments
        for chunk in splitter.split_text(segment["text"])
    ]


def upload_document(
    db,
    vectors,
    data_dir,
    workspace_id,
    name,
    content_or_path: bytes | Path | str,
    file_size: int | None = None,
    content_hash: str | None = None,
):
    """Validate, prevent duplicates, persist the original, then index extracted chunks."""
    name = name.replace("\\", "/").rsplit("/", 1)[-1][:240]
    extension = Path(name).suffix.lower()
    if extension not in SUPPORTED_EXTENSIONS:
        raise ValueError("Supported files: PDF, DOCX, TXT, and Markdown.")

    is_path = isinstance(content_or_path, (str, Path))
    if is_path:
        file_path = Path(content_or_path)
        if file_size is None:
            file_size = file_path.stat().st_size
        if content_hash is None:
            hasher = hashlib.sha256()
            with open(file_path, "rb") as f:
                while chunk := f.read(65536):
                    hasher.update(chunk)
            content_hash = hasher.hexdigest()
    else:
        file_size = len(content_or_path)
        content_hash = hashlib.sha256(content_or_path).hexdigest()

    existing = db.scalar(
        select(Document).where(Document.workspace_id == workspace_id, Document.content_hash == content_hash)
    )
    if existing:
        if is_path:
            Path(content_or_path).unlink(missing_ok=True)
        return existing, True

    log_stage_memory("before_extract_text", doc_name=name)
    segments = extract_text(content_or_path, extension)
    log_stage_memory("after_extract_text", doc_name=name, segment_count=len(segments))

    document = Document(
        id=new_id(),
        workspace_id=workspace_id,
        name=name,
        content_hash=content_hash,
        size=file_size,
        text="\n\n".join(segment["text"] for segment in segments),
    )
    db.add(document)
    try:
        db.commit()
    except IntegrityError:
        db.rollback()
        if is_path:
            Path(content_or_path).unlink(missing_ok=True)
        return db.scalar(
            select(Document).where(
                Document.workspace_id == workspace_id, Document.content_hash == content_hash
            )
        ), True

    try:
        storage.upload_file(workspace_id, document.id, content_or_path)
        index_document(db, vectors, document, segments)
    except Exception as error:
        logging.getLogger("aether").warning(
            "Indexing document %s failed: %s", document.id, type(error).__name__
        )
        document.status = "error"
        document.error = (
            "Indexing failed. Check local model download, disk space, and backend logs; then reindex."
        )
        db.commit()
    finally:
        del segments
        check_and_collect_garbage(threshold_mb=config.gc_threshold_mb)

    return document, False


def index_document(db, vectors, document, segments):
    """Stage a new version. Keep the previous index if rebuilding fails."""
    with vectors.lock:
        version = uuid4().hex
        previous_version = document.index_version
        chunks = chunk_text(segments)
        try:
            vectors.index(document.workspace_id, document.id, document.name, chunks, version)
        except Exception:
            vectors.delete_document(document.workspace_id, document.id, version)
            raise
        document.index_version = version
        document.chunk_count = len(chunks)
        document.status = "ready"
        document.error = None
        try:
            db.commit()
        except Exception:
            db.rollback()
            vectors.delete_document(document.workspace_id, document.id, version)
            raise
        if previous_version:
            vectors.delete_document(document.workspace_id, document.id, previous_version)
        del chunks
        check_and_collect_garbage(threshold_mb=config.gc_threshold_mb)

from io import BytesIO

import pymupdf
import pytest
from docx import Document

from app.services.documents import chunk_text, extract_text


@pytest.mark.parametrize("extension", [".txt", ".md", ".markdown"])
def test_text_extraction(extension):
    segments = extract_text(b"\xef\xbb\xbfHello\x00   world\r\n\r\nResearch", extension)
    assert segments == [{"text": "Hello world\n\nResearch", "page": 0}]


def test_pdf_pages():
    pdf = pymupdf.open()
    for text in ["First research page", "Second research page"]:
        page = pdf.new_page()
        page.insert_text((72, 72), text)
    segments = extract_text(pdf.tobytes(), ".pdf")
    pdf.close()
    assert [item["page"] for item in segments] == [1, 2]
    assert "Second" in segments[1]["text"]


def test_word_paragraphs_and_tables():
    document = Document()
    document.add_paragraph("A useful paragraph")
    document.add_table(rows=1, cols=1).cell(0, 0).text = "Table evidence"
    buffer = BytesIO()
    document.save(buffer)
    text = extract_text(buffer.getvalue(), ".docx")[0]["text"]
    assert "paragraph" in text and "Table evidence" in text


@pytest.mark.parametrize(
    "content,extension",
    [(b"", ".txt"), (b"bad PDF", ".pdf"), (b"bad Word", ".docx"), (b"\xff", ".txt"), (b"hi", ".exe")],
)
def test_invalid_documents(content, extension):
    with pytest.raises(ValueError):
        extract_text(content, extension)


def test_chunk_size_overlap_and_page_metadata():
    text = "a" * 1900
    chunks = chunk_text([{"text": text, "page": 4}])
    assert [len(item["text"]) for item in chunks] == [800, 800, 600]
    assert all(item["page"] == 4 for item in chunks)
    assert chunks[0]["text"][-150:] == chunks[1]["text"][:150]


def test_real_chroma_retrieval_isolation_and_delete(vectors):
    chunks = [
        {"text": "Astronomy studies stars and galaxies", "page": 1},
        {"text": "Bread is made with flour and yeast", "page": 2},
    ]
    vectors.index("workspace_a", "doc_a", "astronomy.txt", chunks, "v1")
    vectors.index("workspace_b", "doc_b", "secret.txt", [{"text": "Private bananas", "page": 0}], "v1")
    results = vectors.retrieve("workspace_a", "stars astronomy", 1)
    assert results[0]["document_id"] == "doc_a"
    assert "Astronomy" in results[0]["text"]
    assert all(item["document_id"] != "doc_b" for item in results)
    vectors.delete_document("workspace_a", "doc_a")
    assert vectors.retrieve("workspace_a", "astronomy") == []
    assert vectors.retrieve("workspace_b", "bananas")


def test_persistent_vectors_after_reopening(vectors, tmp_path):
    from app.services.vector import VectorStore
    from conftest import TestEmbedder

    vectors.index("persistent", "doc", "notes.md", [{"text": "Research evidence", "page": 0}], "v1")
    reopened = VectorStore(tmp_path, "test-only", TestEmbedder())
    assert reopened.retrieve("persistent", "evidence")[0]["document_id"] == "doc"
